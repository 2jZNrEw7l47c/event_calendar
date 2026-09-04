import openpyxl

import failure_log


def test_record_failure_starts_streak_at_one():
    log = {}
    failure_log.record_failure(log, "Casbah", "HTTP 500", "server error",
                                "2026-09-03 10:00")
    assert log["Casbah"]["consecutive_fails"] == 1


def test_record_failure_increments_existing_streak():
    log = {"Casbah": {"error_code": "HTTP 500", "error_message": "server error",
                       "consecutive_fails": 4, "last_failed": "2026-09-02 10:00"}}
    failure_log.record_failure(log, "Casbah", "HTTP 500", "server error",
                                "2026-09-03 10:00")
    assert log["Casbah"]["consecutive_fails"] == 5
    assert log["Casbah"]["last_failed"] == "2026-09-03 10:00"


def test_record_success_clears_entry():
    log = {"Casbah": {"error_code": "HTTP 500", "error_message": "server error",
                       "consecutive_fails": 3, "last_failed": "2026-09-02 10:00"}}
    failure_log.record_success(log, "Casbah")
    assert "Casbah" not in log


def test_record_success_on_absent_venue_is_a_noop():
    log = {}
    failure_log.record_success(log, "Casbah")
    assert log == {}


def test_classify_http_error_uses_status_code():
    class FakeResponse:
        status_code = 404

    class FakeHTTPError(Exception):
        def __init__(self):
            super().__init__("404 Client Error")
            self.response = FakeResponse()

    code, message = failure_log.classify(FakeHTTPError())
    assert code == "HTTP 404"
    assert message == "404 Client Error"


def test_classify_generic_exception_uses_type_name():
    code, message = failure_log.classify(ConnectionError("boom"))
    assert code == "ConnectionError"
    assert message == "boom"


def test_load_missing_file_returns_empty_dict(tmp_path):
    path = str(tmp_path / "failure_log.xlsx")
    assert failure_log.load(path) == {}


def test_save_then_load_round_trip(tmp_path):
    path = str(tmp_path / "failure_log.xlsx")
    log = {"Casbah": {"error_code": "HTTP 500", "error_message": "server error",
                       "consecutive_fails": 2, "last_failed": "2026-09-03 10:00"}}
    failure_log.save(path, log)

    loaded = failure_log.load(path)
    assert loaded == log


def test_save_escapes_formula_injection_in_error_message(tmp_path):
    path = str(tmp_path / "failure_log.xlsx")
    log = {
        "Casbah": {"error_code": "HTTP 500",
                   "error_message": "=cmd|'/c calc'!A1",
                   "consecutive_fails": 1, "last_failed": "2026-09-03 10:00"},
    }
    failure_log.save(path, log)

    wb = openpyxl.load_workbook(path)
    ws = wb["Failures"]
    cell = list(ws.iter_rows(min_row=2, max_row=2))[0][2]
    assert cell.data_type == "s"


def test_save_then_load_round_trip_preserves_formula_like_message(tmp_path):
    path = str(tmp_path / "failure_log.xlsx")
    log = {
        "Casbah": {"error_code": "HTTP 500",
                   "error_message": "=cmd|'/c calc'!A1",
                   "consecutive_fails": 1, "last_failed": "2026-09-03 10:00"},
    }
    failure_log.save(path, log)

    loaded = failure_log.load(path)
    assert loaded == log


def test_load_unreadable_file_prints_warning_and_returns_empty_dict(tmp_path, capsys):
    path = str(tmp_path / "failure_log.xlsx")
    with open(path, "wb") as f:
        f.write(b"not a real xlsx file")

    result = failure_log.load(path)

    captured = capsys.readouterr()
    assert result == {}
    assert "could not read" in captured.out


def test_save_with_illegal_xml_character_does_not_raise(tmp_path, capsys):
    path = str(tmp_path / "failure_log.xlsx")
    log = {
        "Casbah": {"error_code": "HTTP 500", "error_message": "server error\x0bboom",
                   "consecutive_fails": 1, "last_failed": "2026-09-03 10:00"},
    }
    failure_log.save(path, log)

    captured = capsys.readouterr()
    assert "could not write" in captured.out


def test_save_writes_frequent_failures_sheet_filtered_by_threshold(tmp_path):
    path = str(tmp_path / "failure_log.xlsx")
    log = {
        "Casbah": {"error_code": "HTTP 500", "error_message": "server error",
                   "consecutive_fails": 5, "last_failed": "2026-09-03 10:00"},
        "Soda Bar": {"error_code": "Timeout", "error_message": "timed out",
                     "consecutive_fails": 2, "last_failed": "2026-09-03 10:00"},
    }
    failure_log.save(path, log)

    wb = openpyxl.load_workbook(path)
    ws = wb["Frequent Failures"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert [r[0] for r in rows] == ["Casbah"]
